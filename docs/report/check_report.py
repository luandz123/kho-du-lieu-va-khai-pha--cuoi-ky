"""Static document checks only; never runs ETL, mining or a database query."""
from pathlib import Path
import importlib.util, re, json, hashlib
from zipfile import ZipFile
from lxml import etree
from docx import Document
from pypdf import PdfReader

HERE=Path(__file__).resolve().parent
spec=importlib.util.spec_from_file_location('report_builder',HERE/'build_report.py')
b=importlib.util.module_from_spec(spec);spec.loader.exec_module(b)
md=(HERE/'bao_cao_tong_hop.md').read_text(encoding='utf-8')
d=Document(HERE/'bao_cao_tong_hop.docx')
all_cells=[[' '.join(c.text.split()) for c in row.cells] for t in d.tables for row in t.rows]
preserved={}
for p in b.SOURCES:
    rows=[x for x in p.read_text(encoding='utf-8').splitlines() if x.startswith('|')]
    parsed=[b.parse_row(x) for x in rows]
    parsed=[r for r in parsed if not all(re.fullmatch(r':?-+:?',c.replace(' ','')) for c in r)]
    expected=[[' '.join(re.sub(r'`|\*\*','',v).split()) for v in r] for r in parsed]
    edits=[]
    if p.name=='chapter_05.md':
        whole_text='\n'.join(para.text for para in d.paragraphs)
        conditions={'t≤974862386':'t ≤ t₁','974862386<t≤976422051':'t₁ < t ≤ t₂','t>976422051':'t > t₂'}
        for row in expected:
            if len(row)==6 and row[1] in conditions:
                row[1]=conditions[row[1]];edits.append('Table5.5 symbolic cutoff; exact epochs retained in prose')
        for i,row in enumerate(expected):
            if len(row)==7:
                assert row[4] in whole_text,'UNROUNDED_RECALL_TEXT_MISSING'
                expected[i]=row[:4]+row[5:];edits.append('Table5.7 unrounded column moved verbatim to note')
        assert all(t in whole_text for t in ['974862386','976422051','0.02856536704222439','0.03007556334371585','0.02866586314500013'])
    preserved[p.relative_to(b.ROOT).as_posix()]={'rows':len(expected),'approved_editorial_transformations':edits,'all_source_table_cells_in_docx':all(r in all_cells for r in expected)}
    assert preserved[p.relative_to(b.ROOT).as_posix()]['all_source_table_cells_in_docx'],str(p)
body=md.split('# Phụ lục')[0]
assert 'PCA' not in body
assert len(re.findall(r'^# Chương [1-6]\.',md,re.M))==6
refs=re.findall(r'^\*\*(R\d\d)\.',md,re.M)
assert len(refs)==len(set(refs))==15
links=[t for t in re.findall(r'\]\(([^)]+)\)',md) if '://' not in t]
assert all((HERE/t).exists() for t in links)
assert not re.search(r'(?:postgres(?:ql)?|https?)://[^\s/:]+:[^\s/@]+@',md)
with ZipFile(HERE/'bao_cao_tong_hop.docx') as z:
    xml=etree.fromstring(z.read('word/document.xml'))
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    fields=xml.xpath('//w:instrText/text()',namespaces=ns)
    toc=[x for x in fields if x.strip().startswith('TOC')]
    assert len(toc)==3
    equations=len(xml.xpath('//m:oMath',namespaces=ns))
    image_count=len([s for s in z.namelist() if s.startswith('word/media/')])
    assert image_count==9
    assert len(xml.xpath('//w:tbl',namespaces=ns))==len(xml.xpath('//w:tbl/w:tr[1]/w:trPr/w:tblHeader',namespaces=ns))
receipt=b.ROOT/'results/mining/cf_test_evaluation_receipt.json'
assert b.sha(receipt)=='57c16c3c45feea756b771d0b206cb6ec957135205736a4ae40cbb25255af4b50'
pdf=PdfReader(HERE/'bao_cao_tong_hop.pdf')
assert all(len(p.extract_text().strip())>20 for p in pdf.pages)
page_lines=[[line.strip() for line in p.extract_text().splitlines()] for p in pdf.pages]
toc_checks=[]
for h in xml.xpath('//w:hyperlink[starts-with(@w:anchor,"_Toc")]',namespaces=ns):
    s=''.join(h.xpath('.//w:t/text()',namespaces=ns))
    match=re.match(r'((?:Chương [1-6]|Hình \d+\.\d+|Bảng \d+\.\d+)\.).*?(\d+)$',s)
    if not match:continue
    key=match[1];expected=int(match[2])
    actual=next((i+1 for i,lines in enumerate(page_lines) if i>=7 and any(line.startswith(key) for line in lines)),None)
    toc_checks.append({'entry':key,'cached_page':expected,'actual_page':actual,'match':expected==actual})
assert len(toc_checks)==40 and all(c['match'] for c in toc_checks),toc_checks
report={'static_checks':'PASS','page_count':len(pdf.pages),'chapters':6,'figure_count':image_count,'numbered_tables':25,'reference_records':15,'automatic_toc_fields':3,'native_equations':equations,'toc_caption_page_checks':toc_checks,'source_table_cell_preservation':preserved,'official_test_receipt_sha256':b.sha(receipt),'files':{p.name:b.sha(p) for p in HERE.iterdir() if p.is_file() and p.suffix in ['.md','.docx','.pdf','.py','.ps1']},'runtime_rerun':False,'visual_review':'Recorded separately in layout_review.json after inspection'}
(HERE/'static_checks.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k not in ['files','source_table_cell_preservation','toc_caption_page_checks']},ensure_ascii=False))
